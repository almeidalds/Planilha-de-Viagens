import pandas as pd
import unicodedata
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Padrão exato de colunas baseadas nos seus arquivos de modelo
COLUNAS_CHEGADAS = [
    "Tipo", "Missionários", "ID", "Missões", "Países", "Status do Voo", "Origem", "Cia",
    "Voo", "Partida", "Chegada", "Quant.", "Transportes"
]

COLUNAS_PARTIDAS = [
    "Tipo", "Missionários", "ID", "Países", "Status do Voo", "Loc Cia", "Missões",
    "Partida", "Chegada", "Cia", "Voo", "Quant.", "Partida do CTM", "Transportes"
]

def normalizar_texto(texto):
    if pd.isna(texto): 
        return ""
    nfkd = unicodedata.normalize("NFKD", str(texto))
    return "".join([c for c in nfkd if not unicodedata.combining(c)]).strip().upper()

def aplicar_design(df, modo):
    """Filtra e ordena as colunas para o modelo de design escolhido."""
    colunas_alvo = COLUNAS_CHEGADAS if modo == "Chegadas" else COLUNAS_PARTIDAS
    cols_existentes = [c for c in colunas_alvo if c in df.columns]
    return df[cols_existentes].copy()

def aplicar_cores_e_salvar(df, caminho_arquivo, modo):
    """
    Pinta, mescla as células agrupadas e formata a planilha do Excel exatamente nos moldes anexados.
    """
    # Ordena os voos para que os missionários do mesmo voo fiquem juntos
    col_agrupamento = "Chegada" if modo == "Chegadas" else "Partida"
    if "Voo" in df.columns and col_agrupamento in df.columns:
        df = df.sort_values(by=["Voo", col_agrupamento])
        
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = modo
    
    # --- CONFIGURAÇÃO DE FONTES ---
    font_title1 = Font(name='Verdana', size=22, bold=True, color="000000")
    font_title2 = Font(name='Verdana', size=18, bold=True, color="000000")
    font_header = Font(name='Verdana', size=14, bold=True, color="000000")
    font_data = Font(name='Verdana', size=12, bold=False, color="000000")
    
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    
    border_thin = Border(left=Side(style='thin'), right=Side(style='thin'),
                        top=Side(style='thin'), bottom=Side(style='thin'))

    # --- CONFIGURAÇÃO DE CORES ---
    if modo == "Chegadas":
        bg_title = PatternFill(start_color="D9EAD3", end_color="D9EAD3", fill_type="solid") # Verde Claro
        bg_subtitle = PatternFill(start_color="D9EAD3", end_color="D9EAD3", fill_type="solid")
    else:
        bg_title = PatternFill(start_color="AEA2F4", end_color="AEA2F4", fill_type="solid") # Roxo Escuro
        bg_subtitle = PatternFill(start_color="D5D4F8", end_color="D5D4F8", fill_type="solid") # Roxo Claro
    
    num_cols = len(df.columns)
    
    # 1. LINHA DO TÍTULO PRINCIPAL
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=num_cols)
    cell = ws.cell(row=1, column=1, value=f"{modo.upper()} DE MISSIONÁRIOS - CTM BRASIL")
    cell.font = font_title1
    cell.alignment = align_center
    cell.fill = bg_title
    
    # 2. LINHA DE CABEÇALHOS
    for col_idx, col_name in enumerate(df.columns, 1):
        cell = ws.cell(row=2, column=col_idx, value=col_name)
        cell.font = font_header
        cell.alignment = align_center
        cell.fill = bg_title
        cell.border = border_thin
        
    # 3. LINHA DE SUBTÍTULO
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=num_cols)
    cell = ws.cell(row=3, column=1, value="AEROPORTO DE GUARULHOS / CONGONHAS")
    cell.font = font_title2
    cell.alignment = align_center
    cell.fill = bg_subtitle
    cell.border = border_thin
    
    # 4. INSERIR DADOS DOS MISSIONÁRIOS
    start_data_row = 4
    for r_idx, row in enumerate(df.values, start_data_row):
        for c_idx, val in enumerate(row, 1):
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            cell.font = font_data
            cell.border = border_thin
            # Nomes e locais ficam alinhados à esquerda para melhor leitura
            col_name = df.columns[c_idx-1]
            if col_name in ["Missionários", "Missões", "Loc Cia"]:
                cell.alignment = align_left
            else:
                cell.alignment = align_center
                
    # 5. MÁGICA: MESCLAR CÉLULAS DOS MESMOS VOOS (AGRUPAMENTO)
    if "Voo" in df.columns and col_agrupamento in df.columns:
        if modo == "Chegadas":
            cols_to_merge = ["Origem", "Cia", "Voo", "Partida", "Chegada", "Quant.", "Transportes"]
        else:
            cols_to_merge = ["Partida", "Chegada", "Cia", "Voo", "Quant.", "Partida do CTM", "Transportes"]
            
        col_indices_to_merge = [df.columns.get_loc(c) + 1 for c in cols_to_merge if c in df.columns]
        
        status_voo = df["Status do Voo"].astype(str) if "Status do Voo" in df.columns else pd.Series("", index=df.index)
        df['grupo_mescla'] = (
            df['Voo'].astype(str) + "_" + df[col_agrupamento].astype(str)
            + "_" + status_voo
        )
        grupos = []
        indice_inicio = 0
        
        for i in range(1, len(df)):
            if df['grupo_mescla'].iloc[i] != df['grupo_mescla'].iloc[i-1]:
                grupos.append((indice_inicio, i-1))
                indice_inicio = i
        grupos.append((indice_inicio, len(df)-1))
        
        for (inicio, fim) in grupos:
            if inicio != fim:
                excel_start = inicio + start_data_row
                excel_end = fim + start_data_row
                for c_idx in col_indices_to_merge:
                    ws.merge_cells(start_row=excel_start, start_column=c_idx, end_row=excel_end, end_column=c_idx)
                    
    # 6. AJUSTE AUTOMÁTICO DE LARGURA DE COLUNAS (CORRIGIDO)
    for c_idx, col in enumerate(ws.columns, 1):
        max_length = 0
        col_letter = get_column_letter(c_idx) # Pega a letra matematicamente
        for cell in col:
            # Ignora células mescladas (elas não têm largura individual e causam erro)
            if type(cell).__name__ == 'MergedCell':
                continue
                
            # Ignora as linhas de título gigantes (linhas 1 e 3)
            if (cell.row > 3 or cell.row == 2) and cell.value:
                max_length = max(max_length, len(str(cell.value)))
                
        ws.column_dimensions[col_letter].width = min(max_length + 3, 50)
        
    # Trava as três primeiras linhas
    ws.freeze_panes = "A4"
    wb.save(caminho_arquivo)